# Injecte la DMT (Export Vonage Call) dans data/data.json : conserve le reste, remplace les jours/activités présents dans l'export.
# Usage : python scripts/build_dmt.py [Export_Vonage.xlsx] [data.json]
import json, sys, os, re, datetime, warnings, unicodedata
warnings.filterwarnings('ignore')
import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_X = sys.argv[1] if len(sys.argv) > 1 else '/home/z/my-project/upload/Export Vonage Call.xlsx'
SRC_J = sys.argv[2] if len(sys.argv) > 2 else os.path.join(BASE, 'data', 'data.json')

def strip_acc(s):
    return ''.join(ch for ch in unicodedata.normalize('NFD', str(s)) if unicodedata.category(ch) != 'Mn').lower()

def num(v):
    if v is None or v == '': return 0.0
    if isinstance(v, (int, float)): return float(v)
    try: return float(str(v).replace(',', '.'))
    except ValueError: return 0.0

def isod(v):
    if isinstance(v, datetime.datetime): return v.date().isoformat()
    if isinstance(v, datetime.date): return v.isoformat()
    s = str(v).strip()
    m = re.match(r'^(\d{4})[\/\-.](\d{1,2})[\/\-.](\d{1,2})', s)
    if m: return f'{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}'
    m = re.match(r'^(\d{1,2})[\/\-.](\d{1,2})[\/\-.](\d{2,4})$', s)
    if m: return f'{int(m.group(3)):04d}-{int(m.group(2)):02d}-{int(m.group(1)):02d}'
    return ''

wb = openpyxl.load_workbook(SRC_X, data_only=True)
ws = wb.worksheets[0]
heads = {str(ws.cell(1, c).value).strip(): c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}
def find(*subs):
    for k, c in heads.items():
        if all(s in strip_acc(k) for s in subs): return c
    return None
C_COMP, C_REP = find('comp'), find('repondu')
C_DATE = find('date de creation') or find('date')
C_COM, C_ACW = find('duree de communication'), find('after call work')
if not (C_COMP and C_DATE):
    raise SystemExit('Colonnes Compétence / Date de création introuvables')

agg = {}                                   # (date, flux) -> [n répondus, sum com, sum acw]
nrep = 0
for r in range(2, ws.max_row + 1):
    d = isod(ws.cell(r, C_DATE).value)
    if not d: continue
    fl = str(ws.cell(r, C_COMP).value or '').strip().upper()
    if not fl: continue
    rep = num(ws.cell(r, C_REP).value) if C_REP else 1
    if rep <= 0: continue                  # seuls les appels répondus entrent dans la DMT
    com = num(ws.cell(r, C_COM).value) if C_COM else 0
    acw = num(ws.cell(r, C_ACW).value) if C_ACW else 0
    a = agg.setdefault((d, fl), [0, 0.0, 0.0])
    a[0] += 1; a[1] += com; a[2] += acw
    nrep += 1
if not agg:
    raise SystemExit('Aucune ligne exploitable')

data = json.load(open(SRC_J, encoding='utf-8')) if os.path.exists(SRC_J) else {'meta': {}, 'roadmap': [], 'capa': {'dates': [], 'flux': []}}
M = data.get('dmt') or {'file': '', 'maj': '', 'dates': [], 'flows': [], 'f': {}}
dates = list(M['dates']); flows = list(M['flows']); f = M['f']
for (_, fl) in agg:
    if fl not in flows:
        flows.append(fl); f[fl] = {'n': [0] * len(dates), 'com': [0.0] * len(dates), 'acw': [0.0] * len(dates)}
for d in sorted({d for (d, _) in agg}):
    if d not in dates:
        dates.append(d)
        for fl in flows: f[fl] = {k: f[fl][k] + [0] for k in ('n', 'com', 'acw')}
order = sorted(range(len(dates)), key=lambda k: dates[k])
dates = [dates[k] for k in order]
for fl in flows:
    for k in ('n', 'com', 'acw'): f[fl][k] = [f[fl][k][j] for j in order]
for (d, fl), (n, com, acw) in agg.items():
    di = dates.index(d)
    f[fl]['n'][di], f[fl]['com'][di], f[fl]['acw'][di] = int(n), com, acw
rank = lambda fl: (['HYPER', 'C2E', 'AFR', 'DRIVE', 'SAV'].index(fl) if fl in ['HYPER', 'C2E', 'AFR', 'DRIVE', 'SAV'] else 9)
flows.sort(key=lambda fl: (rank(fl), fl))

data['dmt'] = {'file': os.path.basename(SRC_X), 'maj': datetime.date.today().isoformat(),
               'dates': dates, 'flows': flows, 'f': f}
data.setdefault('meta', {})['dmt'] = os.path.basename(SRC_X)
data['meta']['maj'] = datetime.date.today().isoformat()
json.dump(data, open(SRC_J, 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
tot = sum(a[0] for a in agg.values()); dur = sum(a[1] + a[2] for a in agg.values())
print(f'OK -> {SRC_J} | {len(dates)} jours, {len(flows)} activités | {tot} appels répondus | DMT globale {dur / max(tot, 1):.1f} s | {os.path.getsize(SRC_J)} octets')
