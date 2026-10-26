# Régénère data/data.json : conserve roadmap+capa existants, injecte le réalisé depuis un DRM
# Usage : python scripts/build_data.py [chemin/vers/DRM.xlsx] [data.json_existant]
import json, sys, os, datetime, warnings
warnings.filterwarnings('ignore')
import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DRM  = sys.argv[1] if len(sys.argv) > 1 else '/home/z/my-project/upload/102026 - DRM.xlsx'
SRC  = sys.argv[2] if len(sys.argv) > 2 else os.path.join(BASE, 'data', 'data.json')

FLOWS = ['HYPER', 'C2E', 'AFR', 'DRIVE', 'SAV']

def num(v):
    if v is None or v in ('', '*'): return 0.0
    if isinstance(v, (int, float)): return float(v)
    try: return float(str(v).replace(',', '.').replace(' ', ''))
    except ValueError: return 0.0

def parse_drm(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['DATA_ROADMAP'] if 'DATA_ROADMAP' in wb.sheetnames else wb.worksheets[0]
    COL_T = {f: c for f, c in zip(FLOWS, range(4, 9))}     # D-H  Traités
    COL_R = {f: c for f, c in zip(FLOWS, range(12, 17))}   # L-P  Reçus
    COL_S = {f: c for f, c in zip(FLOWS, range(18, 23))}   # R-V  Traités 60s
    dates, sem, jours = [], [], []
    f = {fl: {'t': [], 'r': [], 's': []} for fl in FLOWS}
    seen = set()
    for r in range(4, ws.max_row + 1):
        c3 = ws.cell(r, 3).value
        if isinstance(c3, str) and c3.strip().upper() == 'DATE':
            break                                           # 2e table (QS/QP) -> fin
        if c3 is None or c3 == '': continue
        d = c3.date().isoformat() if isinstance(c3, datetime.datetime) else str(c3)[:10]
        if d in seen: continue
        seen.add(d)
        dates.append(d)
        sem.append(ws.cell(r, 1).value or 0)
        jours.append(str(ws.cell(r, 2).value or '').strip())
        for fl in FLOWS:
            f[fl]['t'].append(num(ws.cell(r, COL_T[fl]).value))
            f[fl]['r'].append(num(ws.cell(r, COL_R[fl]).value))
            f[fl]['s'].append(num(ws.cell(r, COL_S[fl]).value))
    order = sorted(range(len(dates)), key=lambda k: dates[k])
    dates = [dates[k] for k in order]
    sem   = [sem[k] for k in order]
    jours = [jours[k] for k in order]
    for fl in FLOWS:
        for k in ('t', 'r', 's'):
            f[fl][k] = [f[fl][k][j] for j in order]
    return {'maj': datetime.date.today().isoformat(), 'file': os.path.basename(path),
            'flows': FLOWS, 'dates': dates, 'sem': sem, 'jours': jours, 'f': f}

data = json.load(open(SRC, encoding='utf-8')) if os.path.exists(SRC) else {'meta': {}, 'roadmap': [], 'capa': {'dates': [], 'flux': []}}
real = parse_drm(DRM)

# fusion : dates du nouveau DRM remplacées/ajoutées
old = data.get('real') or {'dates': [], 'sem': [], 'jours': [], 'f': {fl: {'t': [], 'r': [], 's': []} for fl in FLOWS}}
keep = [i for i, d in enumerate(old['dates']) if d not in set(real['dates'])]
merged = {'maj': real['maj'], 'file': real['file'], 'flows': FLOWS,
          'dates': [old['dates'][i] for i in keep] + real['dates'],
          'sem':   [old['sem'][i] for i in keep] + real['sem'],
          'jours': [old['jours'][i] for i in keep] + real['jours'], 'f': {}}
for fl in FLOWS:
    merged['f'][fl] = {k: [old['f'][fl][k][i] for i in keep] + real['f'][fl][k] for k in ('t', 'r', 's')}
order = sorted(range(len(merged['dates'])), key=lambda k: merged['dates'][k])
for key in ('dates', 'sem', 'jours'):
    merged[key] = [merged[key][i] for i in order]
for fl in FLOWS:
    for k in ('t', 'r', 's'):
        merged['f'][fl][k] = [merged['f'][fl][k][i] for i in order]

data['real'] = merged
data['meta']['real'] = real['file']
data['meta']['maj'] = datetime.date.today().isoformat()
out = os.path.join(BASE, 'data', 'data.json')
json.dump(data, open(out, 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
print(f'OK -> {out} | {len(merged["dates"])} jours | {os.path.getsize(out)} octets')
